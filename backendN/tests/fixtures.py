"""Builds small, real documents in memory so tests need no binary fixtures."""

from __future__ import annotations

import io
import zipfile


def make_pdf(text: str) -> bytes:
    """Single-page PDF with one line of Helvetica text and a valid xref."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def _zip(entries: dict[str, str | bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in entries.items():
            archive.writestr(name, body)
    return buf.getvalue()


def make_docx(heading: str, paragraph: str) -> bytes:
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    return _zip(
        {
            "[Content_Types].xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/word/document.xml" ContentType="application/'
                'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                "</Types>"
            ),
            "_rels/.rels": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/officeDocument" Target="word/document.xml"/>'
                "</Relationships>"
            ),
            "word/document.xml": (
                f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{w}"><w:body>'
                f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{heading}</w:t></w:r></w:p>'
                f"<w:p><w:r><w:t>{paragraph}</w:t></w:r></w:p>"
                "</w:body></w:document>"
            ),
        }
    )


def make_xlsx(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx(title: str, body: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_epub(chapter_text: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        # EPUB requires an uncompressed "mimetype" entry first.
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        archive.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
            'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            '<dc:title>Test Book</dc:title><dc:identifier id="id">test-book</dc:identifier>'
            "<dc:language>en</dc:language></metadata>"
            '<manifest><item id="c1" href="chapter1.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/></spine></package>',
        )
        archive.writestr(
            "OEBPS/chapter1.xhtml",
            '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Ch 1</title></head>'
            f"<body><h1>Chapter One</h1><p>{chapter_text}</p></body></html>",
        )
    return buf.getvalue()


def make_table_pdf(pages: int, table_pages: set[int]) -> bytes:
    """PDF where the given pages hold a 3-column borderless table (form-like
    to MarkItDown's detector) and the others hold plain text."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    font_id = 3 + 2 * pages
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(pages))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>".encode())
    for i in range(pages):
        ops = ["BT /F1 11 Tf"]
        if i in table_pages:
            ops.append(f"1 0 0 1 50 760 Tm (marker-table-{i}) Tj")
            for r in range(30):
                y = 730 - r * 20
                for x, cell in ((50, f"Item-{r}"), (260, f"{r * 7} units"), (460, f"USD {r * 13}")):
                    ops.append(f"1 0 0 1 {x} {y} Tm ({cell}) Tj")
        else:
            ops.append("14 TL 50 760 Td")
            ops.append(f"(marker-page-{i}) Tj T*")
            ops += [f"(Plain prose line {n} on page {i} about markdown conversion.) Tj T*" for n in range(40)]
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()
